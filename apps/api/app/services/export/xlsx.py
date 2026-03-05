"""XLSX export generator with multi-sheet workbook."""

import logging
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.invoice import Invoice
from app.services.export.core import (
    ACCOUNTING_HEADERS_SR,
    INVOICE_HEADERS_SR,
    LINE_ITEM_HEADERS_SR,
    extract_accounting_row,
    extract_invoice_row,
    format_serbian_number,
)

logger = logging.getLogger(__name__)

# Styling constants
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def generate_xlsx(
    invoices: list[Invoice],
    include_line_items: bool = True,
    date_format: str = "DD.MM.YYYY",
    decimal_separator: str = ",",
) -> BytesIO:
    """Generate multi-sheet XLSX workbook.

    Args:
        invoices: List of Invoice instances (with accounting_intent loaded).
        include_line_items: Whether to include line items sheet.
        date_format: Date format string.
        decimal_separator: Decimal separator for numbers.

    Returns:
        BytesIO buffer containing the XLSX file.
    """
    wb = Workbook()

    # Sheet 1: Invoices
    _build_invoices_sheet(wb.active, invoices, date_format, decimal_separator)

    # Sheet 2: Line Items (optional)
    if include_line_items:
        ws_items = wb.create_sheet("Stavke")
        _build_line_items_sheet(ws_items, invoices, decimal_separator)

    # Sheet 3: Accounting (only if any invoice has accounting_intent)
    has_accounting = any(inv.accounting_intent for inv in invoices)
    if has_accounting:
        ws_acc = wb.create_sheet("Knjizenje")
        _build_accounting_sheet(ws_acc, invoices, decimal_separator)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _style_header_row(ws, headers: list[str]) -> None:
    """Apply header styling to the first row."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _auto_width(ws) -> None:
    """Auto-fit column widths based on content."""
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)


def _build_invoices_sheet(ws, invoices, date_format, decimal_separator):
    """Build the main Invoices sheet."""
    ws.title = "Fakture"
    headers = list(INVOICE_HEADERS_SR.values())
    _style_header_row(ws, headers)

    for row_idx, inv in enumerate(invoices, 2):
        row_data = extract_invoice_row(inv, date_format, decimal_separator)
        for col_idx, key in enumerate(INVOICE_HEADERS_SR.keys(), 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data[key])
            cell.border = THIN_BORDER

    _auto_width(ws)


def _build_line_items_sheet(ws, invoices, decimal_separator):
    """Build the Line Items sheet."""
    headers = list(LINE_ITEM_HEADERS_SR.values())
    _style_header_row(ws, headers)

    row_idx = 2
    for inv in invoices:
        if not inv.line_items or not isinstance(inv.line_items, list):
            continue
        for item in inv.line_items:
            if not isinstance(item, dict):
                continue
            values = [
                inv.invoice_number or "",
                item.get("description", ""),
                format_serbian_number(
                    Decimal(str(item["quantity"])) if item.get("quantity") is not None else None,
                    decimal_separator,
                ),
                format_serbian_number(
                    Decimal(str(item["unit_price"]))
                    if item.get("unit_price") is not None
                    else None,
                    decimal_separator,
                ),
                format_serbian_number(
                    Decimal(str(item["total"])) if item.get("total") is not None else None,
                    decimal_separator,
                ),
                format_serbian_number(
                    Decimal(str(item["tax_rate"])) if item.get("tax_rate") is not None else None,
                    decimal_separator,
                ),
                format_serbian_number(
                    Decimal(str(item["tax_amount"]))
                    if item.get("tax_amount") is not None
                    else None,
                    decimal_separator,
                ),
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = THIN_BORDER
            row_idx += 1

    _auto_width(ws)


def _build_accounting_sheet(ws, invoices, decimal_separator):
    """Build the Accounting sheet."""
    headers = list(ACCOUNTING_HEADERS_SR.values())
    _style_header_row(ws, headers)

    row_idx = 2
    for inv in invoices:
        row_data = extract_accounting_row(inv, decimal_separator)
        if row_data is None:
            continue
        for col_idx, key in enumerate(ACCOUNTING_HEADERS_SR.keys(), 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data[key])
            cell.border = THIN_BORDER
        row_idx += 1

    _auto_width(ws)
