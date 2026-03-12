"""PDV book (KPR/KIR) generation service.

Queries AccountingIntent.pdv_book_entries for a given organization, period,
and book type, then generates XLSX or CSV exports with proper Serbian column
headers, data rows, and totals row. The totals map directly to PP-PDV form
fields for the Serbian VAT return.
"""

from __future__ import annotations

import csv
import logging
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import Integer, cast, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.accounting_intent import AccountingIntent
from app.models.invoice import Invoice
from app.services.export.xlsx import (
    THIN_BORDER,
    _auto_width,
    _style_header_row,
)

logger = logging.getLogger(__name__)

# Column definitions per book type
# Each tuple: (header_label, entry_key, is_numeric)
KPR_COLUMNS: list[tuple[str, str, bool]] = [
    ("R.br.", "sequence", False),
    ("Datum prijema", "entry_date", False),
    ("Datum fakture", "invoice_date", False),
    ("Broj fakture", "invoice_number", False),
    ("Naziv dobavljača", "counterparty_name", False),
    ("PIB dobavljača", "counterparty_pib", False),
    ("Ukupno sa PDV", "total", True),
    ("Osnovica 20%", "base_20", True),
    ("PDV 20%", "vat_20", True),
    ("Osnovica 10%", "base_10", True),
    ("PDV 10%", "vat_10", True),
    ("Polje 8 PP-PDV", "pp_pdv_polje_8", True),
    ("Polje 8a PP-PDV", "pp_pdv_polje_8a", True),
    ("Polje 9 PP-PDV", "pp_pdv_polje_9", True),
]

KIR_COLUMNS: list[tuple[str, str, bool]] = [
    ("R.br.", "sequence", False),
    ("Datum izdavanja", "entry_date", False),
    ("Datum fakture", "invoice_date", False),
    ("Broj fakture", "invoice_number", False),
    ("Naziv kupca", "counterparty_name", False),
    ("PIB kupca", "counterparty_pib", False),
    ("Ukupno sa PDV", "total", True),
    ("Osnovica 20%", "base_20", True),
    ("PDV 20%", "vat_20", True),
    ("Osnovica 10%", "base_10", True),
    ("PDV 10%", "vat_10", True),
    ("Polje 3 PP-PDV", "pp_pdv_polje_3", True),
    ("Polje 4 PP-PDV", "pp_pdv_polje_4", True),
    ("Polje 6 PP-PDV", "pp_pdv_polje_6", True),
    ("Polje 6a PP-PDV", "pp_pdv_polje_6a", True),
]


def _get_columns(book_type: str) -> list[tuple[str, str, bool]]:
    """Get column definitions for a book type.

    Args:
        book_type: KPR or KIR.

    Returns:
        List of (header, key, is_numeric) tuples.
    """
    return KPR_COLUMNS if book_type == "KPR" else KIR_COLUMNS


def _flatten_entry(entry: dict, book_type: str) -> dict[str, str]:
    """Flatten a pdv_book_entries dict, extracting pp_pdv_fields into top-level keys.

    Args:
        entry: Raw pdv_book_entries dict from AccountingIntent.
        book_type: KPR or KIR.

    Returns:
        Flattened dict with pp_pdv_polje_X keys.
    """
    flat = dict(entry)
    pp_pdv = entry.get("pp_pdv_fields", {})

    if book_type == "KPR":
        # Combine base and tax for each polje into single value
        polje_8_base = Decimal(pp_pdv.get("polje_8_1", "0"))
        polje_8_tax = Decimal(pp_pdv.get("polje_8_2", "0"))
        flat["pp_pdv_polje_8"] = str(polje_8_base + polje_8_tax)

        polje_8a_base = Decimal(pp_pdv.get("polje_8a_1", "0"))
        polje_8a_tax = Decimal(pp_pdv.get("polje_8a_2", "0"))
        flat["pp_pdv_polje_8a"] = str(polje_8a_base + polje_8a_tax)

        polje_9_base = Decimal(pp_pdv.get("polje_9_1", "0"))
        polje_9_tax = Decimal(pp_pdv.get("polje_9_2", "0"))
        flat["pp_pdv_polje_9"] = str(polje_9_base + polje_9_tax)
    else:
        flat["pp_pdv_polje_3"] = pp_pdv.get("polje_3_1", "0")
        flat["pp_pdv_polje_4"] = pp_pdv.get("polje_4_1", "0")
        flat["pp_pdv_polje_6"] = pp_pdv.get("polje_6_1", "0")
        flat["pp_pdv_polje_6a"] = pp_pdv.get("polje_6a_1", "0")

    return flat


async def fetch_pdv_book_entries(
    db: AsyncSession,
    organization_id: UUID,
    period: str,
    book_type: str,
    client_id: UUID | None = None,
) -> list[dict]:
    """Query all pdv_book_entries for org/period/book_type, ordered by sequence.

    Args:
        db: Database session.
        organization_id: Organization UUID.
        period: Period string (YYYY-MM).
        book_type: KPR or KIR.
        client_id: Optional client filter for agency users.

    Returns:
        List of flattened pdv_book_entries dicts, sorted by sequence number.
    """
    query = (
        select(AccountingIntent.pdv_book_entries)
        .where(
            AccountingIntent.organization_id == organization_id,
            AccountingIntent.pdv_book_entries["book_type"].as_string() == book_type,
            AccountingIntent.pdv_book_entries["period"].as_string() == period,
        )
        .order_by(
            cast(
                AccountingIntent.pdv_book_entries["sequence"].as_string(),
                Integer,
            )
        )
    )

    if client_id is not None:
        query = query.join(Invoice, AccountingIntent.invoice_id == Invoice.id).where(
            Invoice.client_id == client_id
        )

    result = await db.execute(query)
    raw_entries = result.scalars().all()

    return [_flatten_entry(e, book_type) for e in raw_entries if e]


def generate_pdv_book_xlsx(
    entries: list[dict],
    book_type: str,
    period: str,
) -> BytesIO:
    """Generate XLSX workbook for a KPR or KIR book.

    Creates a single-sheet workbook with Serbian column headers, data rows,
    and a bold totals row at the bottom.

    Args:
        entries: Flattened pdv_book_entries dicts.
        book_type: KPR or KIR.
        period: Period string for the sheet title.

    Returns:
        BytesIO buffer containing the XLSX file.
    """
    columns = _get_columns(book_type)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{book_type} {period}"

    # Header row
    headers = [col[0] for col in columns]
    _style_header_row(ws, headers)

    # Data rows
    for row_idx, entry in enumerate(entries, 2):
        for col_idx, (_, key, is_numeric) in enumerate(columns, 1):
            raw = entry.get(key, "")
            if is_numeric and raw:
                try:
                    value = float(Decimal(str(raw)))
                except (ValueError, TypeError):
                    value = raw
            else:
                value = raw
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER

    # Totals row
    if entries:
        totals_row = len(entries) + 2
        totals = _compute_totals(entries, columns)
        bold_font = Font(bold=True, size=11)

        for col_idx, (header, key, is_numeric) in enumerate(columns, 1):
            if col_idx == 1:
                cell = ws.cell(row=totals_row, column=col_idx, value="UKUPNO")
            elif is_numeric:
                cell = ws.cell(row=totals_row, column=col_idx, value=totals.get(key, 0))
            else:
                cell = ws.cell(row=totals_row, column=col_idx, value="")
            cell.font = bold_font
            cell.border = THIN_BORDER

    _auto_width(ws)

    # Number formatting for numeric columns
    for col_idx, (_, _, is_numeric) in enumerate(columns, 1):
        if is_numeric:
            for row in range(2, len(entries) + 3):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdv_book_csv(
    entries: list[dict],
    book_type: str,
    period: str,
) -> BytesIO:
    """Generate CSV file for a KPR or KIR book.

    Args:
        entries: Flattened pdv_book_entries dicts.
        book_type: KPR or KIR.
        period: Period string (for filename, not included in content).

    Returns:
        BytesIO buffer containing the CSV file (UTF-8 with BOM).
    """
    columns = _get_columns(book_type)
    headers = [col[0] for col in columns]

    output = StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_ALL)
    writer.writerow(headers)

    for entry in entries:
        row = []
        for _, key, is_numeric in columns:
            raw = entry.get(key, "")
            if is_numeric and raw:
                try:
                    # Use comma as decimal separator for Serbian locale
                    val = Decimal(str(raw))
                    row.append(str(val).replace(".", ","))
                except (ValueError, TypeError):
                    row.append(str(raw))
            else:
                row.append(str(raw))
        writer.writerow(row)

    # Totals row
    if entries:
        totals = _compute_totals(entries, columns)
        totals_row = []
        for i, (_, key, is_numeric) in enumerate(columns):
            if i == 0:
                totals_row.append("UKUPNO")
            elif is_numeric:
                val = totals.get(key, Decimal("0"))
                totals_row.append(str(val).replace(".", ","))
            else:
                totals_row.append("")
        writer.writerow(totals_row)

    # Encode as UTF-8 with BOM + sep hint for Excel compatibility
    buffer = BytesIO()
    buffer.write(b"\xef\xbb\xbf")
    buffer.write(b"sep=;\n")
    buffer.write(output.getvalue().encode("utf-8"))
    buffer.seek(0)
    return buffer


def _compute_totals(
    entries: list[dict],
    columns: list[tuple[str, str, bool]],
) -> dict[str, float]:
    """Compute column totals for numeric fields.

    Args:
        entries: List of flattened entry dicts.
        columns: Column definitions.

    Returns:
        Dict of column key to total float value.
    """
    totals: dict[str, Decimal] = {}
    for _, key, is_numeric in columns:
        if not is_numeric:
            continue
        total = Decimal("0")
        for entry in entries:
            raw = entry.get(key, "0")
            try:
                total += Decimal(str(raw))
            except (ValueError, TypeError):
                pass
        totals[key] = total

    return {k: float(v) for k, v in totals.items()}
