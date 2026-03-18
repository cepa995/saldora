"""Unit tests for the audit export service.

Tests cover _build_invoice_register, _build_vat_summary, _build_audit_trail,
_add_documents, and generate_audit_export. S3 calls and DB queries are mocked
throughout — no running infrastructure is required.
"""

import csv
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

from app.services.export.audit import (
    UTF8_BOM,
    _add_documents,
    _build_audit_trail,
    _build_invoice_register,
    _build_vat_summary,
    generate_audit_export,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_invoice(**kwargs):
    """Create a SimpleNamespace invoice for unit testing (no DB required).

    Args:
        **kwargs: Field overrides.

    Returns:
        SimpleNamespace with invoice fields.
    """
    inv = SimpleNamespace(
        id=kwargs.get("id", uuid4()),
        organization_id=kwargs.get("organization_id", uuid4()),
        invoice_number=kwargs.get("invoice_number", "RE-2026-001"),
        invoice_date=kwargs.get("invoice_date", date(2026, 3, 1)),
        due_date=kwargs.get("due_date", date(2026, 3, 31)),
        seller=kwargs.get("seller", {"name": "Dobavljač DOO", "pib": "123456789"}),
        buyer=kwargs.get("buyer", {"name": "Kupac DOO", "pib": "987654321"}),
        subtotal=kwargs.get("subtotal", Decimal("10000.00")),
        tax_rate=kwargs.get("tax_rate", Decimal("20.00")),
        tax_amount=kwargs.get("tax_amount", Decimal("2000.00")),
        total_amount=kwargs.get("total_amount", Decimal("12000.00")),
        currency=kwargs.get("currency", "RSD"),
        status=kwargs.get("status", "verified"),
        document_path=kwargs.get("document_path", None),
        document_content_type=kwargs.get("document_content_type", "application/pdf"),
    )
    return inv


def _make_audit_log(**kwargs):
    """Create a SimpleNamespace audit log entry for unit testing.

    Args:
        **kwargs: Field overrides.

    Returns:
        SimpleNamespace with audit log fields.
    """
    return SimpleNamespace(
        id=kwargs.get("id", uuid4()),
        organization_id=kwargs.get("organization_id", uuid4()),
        user_id=kwargs.get("user_id", uuid4()),
        action=kwargs.get("action", "invoice.verified"),
        entity_type=kwargs.get("entity_type", "invoice"),
        entity_id=kwargs.get("entity_id", uuid4()),
        ip_address=kwargs.get("ip_address", "192.168.1.1"),
        created_at=kwargs.get("created_at", datetime(2026, 3, 1, 10, 0, 0, tzinfo=UTC)),
    )


def _parse_csv_bytes(data: bytes, delimiter: str = ";") -> list[list[str]]:
    """Strip the UTF-8 BOM and parse CSV bytes.

    Args:
        data: Raw CSV bytes, possibly with BOM prefix.
        delimiter: Field delimiter character.

    Returns:
        List of rows, each a list of field strings.
    """
    text = data.lstrip(b"\xef\xbb\xbf").decode("utf-8")
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    return list(reader)


# ===========================================================================
# A. _build_invoice_register()
# ===========================================================================


def test_build_invoice_register_has_bom():
    """_build_invoice_register() output starts with a UTF-8 BOM for Excel compatibility."""
    result = _build_invoice_register([_make_invoice()])
    assert result[:3] == UTF8_BOM


def test_build_invoice_register_header_row():
    """_build_invoice_register() first data row is the Serbian column header."""
    result = _build_invoice_register([])
    rows = _parse_csv_bytes(result)
    assert rows[0][0] == "Broj fakture"
    assert "Datum fakture" in rows[0]
    assert "PIB prodavca" in rows[0]
    assert "Ukupan iznos" in rows[0]


def test_build_invoice_register_single_invoice_row():
    """Each invoice produces exactly one data row below the header."""
    inv = _make_invoice(invoice_number="INV-001")
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    # Row 0 = header, row 1 = first invoice
    assert len(rows) == 2
    assert rows[1][0] == "INV-001"


def test_build_invoice_register_seller_fields():
    """Seller name and PIB are written to the correct columns."""
    inv = _make_invoice(seller={"name": "Test Dobavljač", "pib": "111222333"})
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    data_row = rows[1]
    assert "Test Dobavljač" in data_row
    assert "111222333" in data_row


def test_build_invoice_register_buyer_fields():
    """Buyer name and PIB are written to the correct columns."""
    inv = _make_invoice(buyer={"name": "Test Kupac", "pib": "444555666"})
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    data_row = rows[1]
    assert "Test Kupac" in data_row
    assert "444555666" in data_row


def test_build_invoice_register_currency_defaults_to_rsd():
    """Currency defaults to 'RSD' when the invoice currency field is None."""
    inv = _make_invoice()
    inv.currency = None
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    assert "RSD" in rows[1]


def test_build_invoice_register_none_invoice_number():
    """None invoice_number produces an empty string in the output row."""
    inv = _make_invoice()
    inv.invoice_number = None
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    assert rows[1][0] == ""


def test_build_invoice_register_non_dict_seller_buyer():
    """Non-dict seller/buyer objects are treated as empty dicts (no crash)."""
    inv = _make_invoice()
    inv.seller = None
    inv.buyer = None
    result = _build_invoice_register([inv])
    rows = _parse_csv_bytes(result)
    # Seller name and PIB columns should be empty
    assert rows[1][3] == ""
    assert rows[1][4] == ""


def test_build_invoice_register_multiple_invoices():
    """Multiple invoices each produce one row, in input order."""
    inv1 = _make_invoice(invoice_number="INV-001")
    inv2 = _make_invoice(invoice_number="INV-002")
    inv3 = _make_invoice(invoice_number="INV-003")
    result = _build_invoice_register([inv1, inv2, inv3])
    rows = _parse_csv_bytes(result)
    assert len(rows) == 4  # header + 3 data rows
    assert rows[1][0] == "INV-001"
    assert rows[2][0] == "INV-002"
    assert rows[3][0] == "INV-003"


def test_build_invoice_register_empty_list():
    """Empty invoice list produces only the header row."""
    result = _build_invoice_register([])
    rows = _parse_csv_bytes(result)
    assert len(rows) == 1


# ===========================================================================
# B. _build_vat_summary()
# ===========================================================================


def test_build_vat_summary_returns_bytes():
    """_build_vat_summary() returns non-empty bytes (XLSX format)."""
    result = _build_vat_summary([_make_invoice()])
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_build_vat_summary_valid_xlsx():
    """_build_vat_summary() output can be loaded as a valid openpyxl Workbook."""
    from openpyxl import load_workbook

    result = _build_vat_summary([_make_invoice()])
    wb = load_workbook(BytesIO(result))
    assert "PDV pregled" in wb.sheetnames


def test_build_vat_summary_header_row():
    """The first row of the XLSX sheet contains the expected Serbian column headers."""
    from openpyxl import load_workbook

    result = _build_vat_summary([])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]
    header = [ws.cell(1, c).value for c in range(1, 6)]
    assert header[0] == "Stopa PDV (%)"
    assert header[1] == "Broj faktura"
    assert header[2] == "Osnovica"


def test_build_vat_summary_groups_by_tax_rate():
    """Invoices with the same tax rate are aggregated into a single row."""
    from openpyxl import load_workbook

    inv1 = _make_invoice(
        tax_rate=Decimal("20"),
        tax_amount=Decimal("2000"),
        subtotal=Decimal("10000"),
        total_amount=Decimal("12000"),
    )
    inv2 = _make_invoice(
        tax_rate=Decimal("20"),
        tax_amount=Decimal("1000"),
        subtotal=Decimal("5000"),
        total_amount=Decimal("6000"),
    )
    result = _build_vat_summary([inv1, inv2])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]

    # Row 1 = header, row 2 = rate 20 group, row 3 = UKUPNO
    assert ws.cell(2, 1).value == "20"
    assert ws.cell(2, 2).value == 2  # count


def test_build_vat_summary_different_rates_produce_separate_rows():
    """Invoices with different tax rates produce one row each."""
    from openpyxl import load_workbook

    inv_20 = _make_invoice(tax_rate=Decimal("20"))
    inv_10 = _make_invoice(tax_rate=Decimal("10"))
    result = _build_vat_summary([inv_20, inv_10])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]

    rates = {ws.cell(r, 1).value for r in range(2, ws.max_row)}
    assert "20" in rates
    assert "10" in rates


def test_build_vat_summary_totals_row():
    """The last row is the UKUPNO totals row."""
    from openpyxl import load_workbook

    inv = _make_invoice(tax_rate=Decimal("20"))
    result = _build_vat_summary([inv])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]

    last_row = ws.max_row
    assert ws.cell(last_row, 1).value == "UKUPNO"


def test_build_vat_summary_empty_list():
    """_build_vat_summary() with an empty invoice list produces only header + UKUPNO."""
    from openpyxl import load_workbook

    result = _build_vat_summary([])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]

    # With no rates, only header (row 1) and totals row
    assert ws.max_row == 2
    assert ws.cell(2, 1).value == "UKUPNO"


def test_build_vat_summary_none_amounts_treated_as_zero():
    """None subtotal/tax_amount/total_amount values are treated as 0.0."""
    from openpyxl import load_workbook

    inv = _make_invoice()
    inv.subtotal = None
    inv.tax_amount = None
    inv.total_amount = None
    inv.tax_rate = Decimal("20")

    # Should not raise
    result = _build_vat_summary([inv])
    wb = load_workbook(BytesIO(result))
    ws = wb["PDV pregled"]
    # Count for rate 20 should be 1
    assert ws.cell(2, 2).value == 1


# ===========================================================================
# C. _build_audit_trail()
# ===========================================================================


async def test_build_audit_trail_has_bom():
    """_build_audit_trail() output starts with a UTF-8 BOM."""
    log = _make_audit_log()
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [log]
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    assert result[:3] == UTF8_BOM


async def test_build_audit_trail_header_row():
    """_build_audit_trail() first row is the Serbian column header."""
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = []
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert rows[0][0] == "Datum i vreme"
    assert "Akcija" in rows[0]
    assert "IP adresa" in rows[0]


async def test_build_audit_trail_single_log_row():
    """Each audit log entry produces one row below the header."""
    log = _make_audit_log(action="invoice.verified", ip_address="10.0.0.1")
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [log]
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert len(rows) == 2
    assert rows[1][1] == "invoice.verified"
    assert rows[1][5] == "10.0.0.1"


async def test_build_audit_trail_none_user_id():
    """_build_audit_trail() writes empty string when user_id is None."""
    log = _make_audit_log()
    log.user_id = None
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [log]
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert rows[1][2] == ""


async def test_build_audit_trail_none_entity_id():
    """_build_audit_trail() writes empty string when entity_id is None."""
    log = _make_audit_log()
    log.entity_id = None
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [log]
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert rows[1][4] == ""


async def test_build_audit_trail_none_created_at():
    """_build_audit_trail() writes empty string when created_at is None."""
    log = _make_audit_log()
    log.created_at = None
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [log]
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert rows[1][0] == ""


async def test_build_audit_trail_multiple_logs():
    """Multiple audit log entries each produce their own row, in order."""
    logs = [
        _make_audit_log(action="login", ip_address="1.1.1.1"),
        _make_audit_log(action="invoice.created", ip_address="2.2.2.2"),
        _make_audit_log(action="invoice.exported", ip_address="3.3.3.3"),
    ]
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = logs
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_result)

    result = await _build_audit_trail(mock_db, uuid4(), date(2026, 1, 1), date(2026, 12, 31))
    rows = _parse_csv_bytes(result)
    assert len(rows) == 4  # header + 3 logs
    assert rows[1][1] == "login"
    assert rows[2][1] == "invoice.created"
    assert rows[3][1] == "invoice.exported"


# ===========================================================================
# D. _add_documents()
# ===========================================================================


async def test_add_documents_skips_invoices_without_document_path():
    """_add_documents() skips invoices where document_path is None."""
    inv = _make_invoice(document_path=None)
    zf = MagicMock()
    mock_s3 = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        await _add_documents(zf, [inv])

    mock_s3.get_object.assert_not_called()
    zf.writestr.assert_not_called()


async def test_add_documents_writes_to_zip():
    """_add_documents() downloads from S3 and writes each document to the ZIP."""
    inv = _make_invoice(
        invoice_number="INV-001",
        document_path="organizations/x/invoices/y/original.pdf",
        document_content_type="application/pdf",
    )

    body_mock = MagicMock()
    body_mock.read.return_value = b"%PDF-1.4 content"
    mock_s3 = MagicMock()
    mock_s3.get_object.return_value = {"Body": body_mock}

    zf = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        await _add_documents(zf, [inv])

    zf.writestr.assert_called_once()
    call_args = zf.writestr.call_args[0]
    assert call_args[0].startswith("dokumenti/")
    assert call_args[0].endswith(".pdf")
    assert call_args[1] == b"%PDF-1.4 content"


async def test_add_documents_sanitizes_filename():
    """_add_documents() replaces special characters in invoice_number with underscores."""
    inv = _make_invoice(
        invoice_number="INV/2026\\001",
        document_path="org/inv/original.pdf",
        document_content_type="application/pdf",
    )

    body_mock = MagicMock()
    body_mock.read.return_value = b"pdf bytes"
    mock_s3 = MagicMock()
    mock_s3.get_object.return_value = {"Body": body_mock}

    zf = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        await _add_documents(zf, [inv])

    call_args = zf.writestr.call_args[0]
    filename = call_args[0]
    # Slashes and backslashes must be replaced
    assert "/" not in filename.split("dokumenti/")[1]
    assert "\\" not in filename


async def test_add_documents_uses_invoice_id_when_no_number():
    """_add_documents() falls back to invoice ID when invoice_number is None."""
    inv_id = uuid4()
    inv = _make_invoice(
        id=inv_id,
        invoice_number=None,
        document_path="org/inv/original.pdf",
        document_content_type="application/pdf",
    )

    body_mock = MagicMock()
    body_mock.read.return_value = b"pdf"
    mock_s3 = MagicMock()
    mock_s3.get_object.return_value = {"Body": body_mock}

    zf = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        await _add_documents(zf, [inv])

    call_args = zf.writestr.call_args[0]
    assert str(inv_id).replace("-", "_") in call_args[0] or str(inv_id) in call_args[0]


async def test_add_documents_continues_on_s3_error():
    """_add_documents() continues processing remaining invoices if one S3 fetch fails."""
    inv1 = _make_invoice(
        invoice_number="FAIL-001",
        document_path="org/inv/fail.pdf",
        document_content_type="application/pdf",
    )
    inv2 = _make_invoice(
        invoice_number="OK-002",
        document_path="org/inv/ok.pdf",
        document_content_type="application/pdf",
    )

    body_mock = MagicMock()
    body_mock.read.return_value = b"ok content"
    mock_s3 = MagicMock()

    def _get_object(**kwargs):
        if "fail" in kwargs.get("Key", ""):
            raise RuntimeError("S3 error")
        return {"Body": body_mock}

    mock_s3.get_object.side_effect = _get_object

    zf = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        # Should not raise
        await _add_documents(zf, [inv1, inv2])

    # Only the second invoice should be written
    zf.writestr.assert_called_once()
    call_args = zf.writestr.call_args[0]
    assert "OK_002" in call_args[0] or "OK-002" in call_args[0]


async def test_add_documents_uses_pdf_extension_for_unknown_content_type():
    """_add_documents() falls back to 'pdf' extension when content_type is None."""
    inv = _make_invoice(
        invoice_number="INV-EXT",
        document_path="org/inv/file",
        document_content_type=None,
    )

    body_mock = MagicMock()
    body_mock.read.return_value = b"data"
    mock_s3 = MagicMock()
    mock_s3.get_object.return_value = {"Body": body_mock}

    zf = MagicMock()

    with patch("app.services.export.audit.get_s3_client", return_value=mock_s3):
        await _add_documents(zf, [inv])

    call_args = zf.writestr.call_args[0]
    assert call_args[0].endswith(".pdf")


# ===========================================================================
# E. generate_audit_export()
# ===========================================================================


async def test_generate_audit_export_returns_expected_keys():
    """generate_audit_export() returns a dict with all required keys."""
    org_id = uuid4()
    inv = _make_invoice(organization_id=org_id)

    # Mock DB to return the invoice
    mock_scalars = MagicMock()
    mock_scalars.all.return_value = [inv]
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()
    mock_s3_client.put_object.return_value = {}

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch(
            "app.services.export.audit.get_presigned_url",
            return_value="https://s3.example.com/export.zip",
        ),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"audit")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
        )

    assert "download_url" in result
    assert "file_size" in result
    assert "invoice_count" in result
    assert "period" in result
    assert "s3_key" in result


async def test_generate_audit_export_correct_invoice_count():
    """generate_audit_export() invoice_count matches the number of invoices fetched."""
    org_id = uuid4()
    invoices = [_make_invoice(organization_id=org_id) for _ in range(3)]

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = invoices
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()
    mock_s3_client.put_object.return_value = {}

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch(
            "app.services.export.audit.get_presigned_url",
            return_value="https://s3.example.com/export.zip",
        ),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"audit")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
        )

    assert result["invoice_count"] == 3


async def test_generate_audit_export_s3_key_contains_org_and_dates():
    """generate_audit_export() generates an S3 key with org_id and date range."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = []
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch(
            "app.services.export.audit.get_presigned_url",
            return_value="https://s3.example.com/export.zip",
        ),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
        )

    assert str(org_id) in result["s3_key"]
    assert "2026-01-01" in result["s3_key"]
    assert "2026-03-31" in result["s3_key"]


async def test_generate_audit_export_period_in_result():
    """generate_audit_export() result period contains from/to ISO strings."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = []
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch(
            "app.services.export.audit.get_presigned_url", return_value="https://example.com/zip"
        ),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 2, 1),
            date_to=date(2026, 2, 28),
        )

    assert result["period"]["from"] == "2026-02-01"
    assert result["period"]["to"] == "2026-02-28"


async def test_generate_audit_export_download_url_from_presigned():
    """generate_audit_export() download_url comes from get_presigned_url."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = []
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch(
            "app.services.export.audit.get_presigned_url",
            return_value="https://presigned.url/file.zip",
        ),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
        )

    assert result["download_url"] == "https://presigned.url/file.zip"


async def test_generate_audit_export_skips_vat_summary_when_disabled():
    """generate_audit_export() does not include pdv_pregled.xlsx when include_vat_summary=False."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = [_make_invoice()]
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()
    captured_zip_bytes = {}

    def _capture_put(**kwargs):
        captured_zip_bytes["data"] = kwargs["Body"]

    mock_s3_client.put_object.side_effect = _capture_put

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch("app.services.export.audit.get_presigned_url", return_value="https://example.com"),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
            include_vat_summary=False,
        )

    zf = zipfile.ZipFile(BytesIO(captured_zip_bytes["data"]))
    assert "pdv_pregled.xlsx" not in zf.namelist()


async def test_generate_audit_export_skips_audit_trail_when_disabled():
    """Skips revizijski_trag.csv when include_audit_trail=False."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = [_make_invoice()]
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()
    captured_zip_bytes = {}

    def _capture_put(**kwargs):
        captured_zip_bytes["data"] = kwargs["Body"]

    mock_s3_client.put_object.side_effect = _capture_put

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch("app.services.export.audit.get_presigned_url", return_value="https://example.com"),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
            include_audit_trail=False,
        )

    zf = zipfile.ZipFile(BytesIO(captured_zip_bytes["data"]))
    assert "revizijski_trag.csv" not in zf.namelist()


async def test_generate_audit_export_always_includes_register_csv():
    """generate_audit_export() always includes registar_faktura.csv in the ZIP."""
    org_id = uuid4()

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = []
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()
    captured_zip_bytes = {}

    def _capture_put(**kwargs):
        captured_zip_bytes["data"] = kwargs["Body"]

    mock_s3_client.put_object.side_effect = _capture_put

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch("app.services.export.audit.get_presigned_url", return_value="https://example.com"),
        patch("app.services.export.audit._build_audit_trail", new=AsyncMock(return_value=b"")),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
            include_vat_summary=False,
            include_audit_trail=False,
            include_documents=False,
        )

    zf = zipfile.ZipFile(BytesIO(captured_zip_bytes["data"]))
    assert "registar_faktura.csv" in zf.namelist()


async def test_generate_audit_export_file_size_is_positive():
    """generate_audit_export() returns a positive file_size for a non-trivial export."""
    org_id = uuid4()
    invoices = [_make_invoice(organization_id=org_id) for _ in range(2)]

    mock_scalars = MagicMock()
    mock_scalars.all.return_value = invoices
    mock_execute_result = MagicMock()
    mock_execute_result.scalars.return_value = mock_scalars
    mock_db = AsyncMock()
    mock_db.execute = AsyncMock(return_value=mock_execute_result)

    mock_s3_client = MagicMock()

    with (
        patch("app.services.export.audit.get_s3_client", return_value=mock_s3_client),
        patch("app.services.export.audit.get_presigned_url", return_value="https://example.com"),
        patch(
            "app.services.export.audit._build_audit_trail",
            new=AsyncMock(return_value=b"audit data"),
        ),
        patch("app.services.export.audit._add_documents", new=AsyncMock()),
    ):
        result = await generate_audit_export(
            db=mock_db,
            organization_id=org_id,
            date_from=date(2026, 1, 1),
            date_to=date(2026, 3, 31),
        )

    assert result["file_size"] > 0
